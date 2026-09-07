"""Asignacion de personas a cuentas/clientes, desde la planilla de Azerta.

BUK tiene un campo "Cuentas" en custom_attributes pero esta vacio en los 98
empleados, asi que la planilla es la unica fuente. Se cruza con BUK por RUT.

El RUT se usa solo como llave: no se guarda en el directorio ni sale en ninguna
respuesta. De la planilla tampoco se leen las horas por semana, que son
informacion contractual y nadie pregunta por ellas en el chat.
"""

import logging
import re
import unicodedata
from pathlib import Path

from django.conf import settings
from django.core.cache import cache

logger = logging.getLogger(__name__)

HOJA = "Detalle Cuenta-Persona"
COL_CUENTA, COL_PERSONA, COL_RUT = 0, 1, 3
PRIMERA_FILA = 4  # las tres primeras son titulo y encabezado


def normalizar_rut(valor):
    return str(valor or "").replace(".", "").replace("-", "").strip().lower()


def _clave(texto):
    """Para comparar nombres de cuenta sin importar tildes ni mayusculas."""
    plano = unicodedata.normalize("NFKD", str(texto or "").lower())
    plano = "".join(c for c in plano if unicodedata.category(c) != "Mn")
    return re.sub(r"[^a-z0-9]+", " ", plano).strip()


def archivo():
    carpeta = Path(settings.DOCUMENTOS_DIR)
    if not carpeta.exists():
        return None
    hallazgos = sorted(carpeta.glob("*.xlsx"))
    return hallazgos[0] if hallazgos else None


def cargar(forzar=False):
    """Devuelve {clave_cuenta: {"nombre": str, "ruts": set}}.

    Cacheado por fecha de modificacion: reemplazar la planilla actualiza las
    cuentas sin reiniciar nada.
    """
    ruta = archivo()
    if ruta is None:
        return {}

    firma = (str(ruta), ruta.stat().st_mtime_ns)
    if not forzar:
        cacheado = cache.get("cuentas:mapa")
        if cacheado is not None and cacheado.get("firma") == firma:
            return cacheado["cuentas"]

    try:
        import openpyxl
    except ImportError:
        logger.warning("openpyxl no esta instalado: no se leen las cuentas")
        return {}

    try:
        libro = openpyxl.load_workbook(ruta, data_only=True, read_only=True)
        hoja = libro[HOJA] if HOJA in libro.sheetnames else libro[libro.sheetnames[0]]
        filas = list(hoja.iter_rows(min_row=PRIMERA_FILA, values_only=True))
        libro.close()
    except Exception as error:
        logger.warning("no se pudo leer %s: %s", ruta.name, error)
        return {}

    mapa = {}
    for fila in filas:
        if len(fila) <= COL_RUT:
            continue
        cuenta, rut = fila[COL_CUENTA], normalizar_rut(fila[COL_RUT])
        if not cuenta or not rut:
            continue
        clave = _clave(cuenta)
        if not clave:
            continue
        entrada = mapa.setdefault(clave, {"nombre": str(cuenta).strip(), "ruts": set()})
        entrada["ruts"].add(rut)

    cache.set("cuentas:mapa", {"firma": firma, "cuentas": mapa},
              settings.BUK_CACHE_TTL)
    return mapa


def nombres():
    """Nombres de cuenta tal como se escriben en la planilla."""
    return sorted(v["nombre"] for v in cargar().values())


def buscar(texto):
    """Cuenta mencionada en el texto, o None.

    Se prefiere la coincidencia mas larga: "afp capital" antes que "afp".
    """
    plano = f" {_clave(texto)} "
    mejor = None
    for clave, datos in cargar().items():
        if len(clave) >= 3 and f" {clave} " in plano:
            if mejor is None or len(clave) > len(mejor[0]):
                mejor = (clave, datos)
    return mejor[1] if mejor else None


def asignar(directorio):
    """Agrega la lista de cuentas a cada persona del directorio.

    Cruza por RUT y despues lo descarta: la llave no tiene por que quedar
    guardada junto a los datos que si se muestran.
    """
    mapa = cargar()
    por_rut = {}
    for datos in mapa.values():
        for rut in datos["ruts"]:
            por_rut.setdefault(rut, []).append(datos["nombre"])

    for persona in directorio.values():
        persona["cuentas"] = sorted(por_rut.get(persona.pop("_rut", ""), []))
    return directorio

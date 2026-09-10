"""Texto plano desde un .docx subido a Drive.

Drive solo puede exportar los formatos nativos de Google; un .docx que alguien
subio sin convertir se baja tal cual y hay que abrirlo aca. Los encabezados con
estilo "Heading N" se emiten como titulos Markdown para que `chat/documentos.py`
los use para partir el documento en secciones, igual que con un .md.
"""

import logging

logger = logging.getLogger(__name__)


def extraer(ruta):
    try:
        from docx import Document
    except ImportError:
        logger.warning("python-docx no esta instalado: %s se ignora", ruta)
        return ""

    try:
        documento = Document(str(ruta))
    except Exception as error:  # noqa: BLE001
        logger.warning("no pude leer %s: %s", ruta, error)
        return ""

    lineas = []
    for parrafo in documento.paragraphs:
        texto = parrafo.text.strip()
        if not texto:
            continue
        estilo = (parrafo.style.name or "").lower() if parrafo.style else ""
        if estilo.startswith("heading"):
            nivel = "".join(c for c in estilo if c.isdigit()) or "1"
            lineas.append("#" * min(int(nivel), 6) + " " + texto)
        else:
            lineas.append(texto)

    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells]
            if any(celdas):
                lineas.append(" | ".join(celdas))

    return "\n".join(lineas)

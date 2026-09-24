"""Azerta Finder: numero de contacto de cada persona, desde un archivo de
Drive (un .xlsx) compartido DIRECTO con la cuenta de servicio -no vive
dentro de la carpeta que chat/drive.py sincroniza, asi que se lee por su id,
no por nombre dentro de esa carpeta ni via DRIVE_HOJAS_PERMITIDAS.

Acceso restringido por lista de correos (settings.AZERTA_FINDER_USUARIOS),
mismo criterio que chat/salas.py::usuario_habilitado: es informacion de
contacto de terceros, confidencial, y se abre persona por persona, sin
excepcion por rol -ni gerencia ni un superusuario entran si no estan en la
lista.
"""

import io
import logging
import re

from django.conf import settings
from django.core.cache import cache

from .intents import normalizar

logger = logging.getLogger(__name__)

CACHE_FILAS = "finder:filas:v2"

# Encabezados posibles para cada columna, buscados como substring del
# encabezado ya normalizado (sin tildes, en minuscula): la planilla no la
# administra este proyecto, asi que no se asume el nombre exacto de cada
# columna. Nombre y telefono son obligatorias (sin ellas no hay Finder);
# cargo/organizacion/mail son opcionales, para la tarjeta de contacto -si
# la planilla no las tiene, la tarjeta sale con menos datos, no falla.
_CLAVES_NOMBRE = ("nombre",)
_CLAVES_TELEFONO = ("telefono", "fono", "celular", "movil", "whatsapp", "contacto")
_CLAVES_CARGO = ("cargo",)
_CLAVES_ORGANIZACION = ("organizacion", "empresa")
_CLAVES_MAIL = ("mail", "correo", "email")

# (clave del resultado, encabezados que la identifican, si es obligatoria)
_COLUMNAS = (
    ("nombre", _CLAVES_NOMBRE, True),
    ("telefono", _CLAVES_TELEFONO, True),
    ("cargo", _CLAVES_CARGO, False),
    ("organizacion", _CLAVES_ORGANIZACION, False),
    ("mail", _CLAVES_MAIL, False),
)


class FinderError(Exception):
    """Error al leer la planilla, con mensaje listo para el usuario."""


def usuario_habilitado(correo):
    """Si `correo` puede usar Azerta Finder: lista explicita de correos, sin
    excepcion por rol (ver el criterio en el docstring del modulo)."""
    permitidos = getattr(settings, "AZERTA_FINDER_USUARIOS", set())
    return (correo or "").strip().lower() in permitidos


def _columna(encabezados, claves):
    for i, encabezado in enumerate(encabezados):
        plano = normalizar(str(encabezado or ""))
        if any(clave in plano for clave in claves):
            return i
    return None


def _valor(fila, idx):
    if idx is None or idx >= len(fila) or fila[idx] is None:
        return ""
    valor = fila[idx]
    # Un telefono como "56999813647" en la celda, sin formato de texto, Excel
    # lo guarda como numero: openpyxl lo entrega como float (56999813647.0) y
    # sale con el ".0" pegado. Si es un entero exacto, se muestra como tal.
    if isinstance(valor, float) and valor.is_integer():
        valor = int(valor)
    return str(valor).strip()


def _formatear_telefono(valor):
    """"56999813647" o "999813647" -> "+56 9 9981 3647" (celular chileno).

    Solo formatea cuando el valor es EXACTAMENTE un celular chileno (con o
    sin el 56 adelante): la planilla tambien trae fijos y mas de un numero
    separados por "/", y forzarles el mismo formato los dejaria peor -mejor
    mostrarlos tal cual la persona los cargo.
    """
    solo_digitos = re.sub(r"\D", "", valor)
    if len(solo_digitos) == 9 and solo_digitos.startswith("9"):
        solo_digitos = "56" + solo_digitos
    if len(solo_digitos) == 11 and solo_digitos.startswith("569"):
        return f"+{solo_digitos[:2]} {solo_digitos[2]} {solo_digitos[3:7]} {solo_digitos[7:]}"
    return valor


def cargar(forzar=False):
    """Filas de la planilla: [{"nombre", "telefono", "cargo"?, "organizacion"?,
    "mail"?}, ...] -las tres ultimas solo si la planilla tiene esa columna.

    Cacheadas con TTL fijo (settings.AZERTA_FINDER_CACHE_TTL): a diferencia
    de la carpeta sincronizada, no hay un archivo local cuyo mtime avise que
    cambio, asi que no vale la pena distinguir "cache" de "descarga de
    verdad" -se vuelve a descargar entero cuando vence, nada mas. Lanza
    FinderError si el archivo no esta configurado o no se pudo leer, para que
    la herramienta que llama pueda avisar en vez de decir "no encontrada" por
    un problema que no tiene nada que ver con el nombre buscado.
    """
    if not forzar:
        cacheado = cache.get(CACHE_FILAS)
        if cacheado is not None:
            return cacheado

    file_id = getattr(settings, "AZERTA_FINDER_FILE_ID", "")
    if not file_id or not settings.GOOGLE_DRIVE_CREDENTIALS:
        raise FinderError("Azerta Finder no está disponible todavía. Avisa al equipo técnico.")

    from . import drive

    try:
        contenido = drive.descargar_archivo(file_id)
    except Exception as error:
        logger.warning("azerta finder: no pude descargar la planilla: %s", error)
        raise FinderError("No pude leer la planilla de contactos en este momento.")

    try:
        import openpyxl
    except ImportError:
        raise FinderError("Azerta Finder no está disponible todavía. Avisa al equipo técnico.")

    try:
        libro = openpyxl.load_workbook(io.BytesIO(contenido), data_only=True, read_only=True)
        hoja = libro[libro.sheetnames[0]]
        crudas = list(hoja.iter_rows(values_only=True))
        libro.close()
    except Exception as error:
        logger.warning("azerta finder: no se pudo leer la planilla: %s", error)
        raise FinderError("No pude leer la planilla de contactos en este momento.")

    if not crudas:
        filas = []
    else:
        indices = {clave: _columna(crudas[0], claves) for clave, claves, _ in _COLUMNAS}
        faltan_obligatorias = [clave for clave, _, obligatoria in _COLUMNAS
                               if obligatoria and indices[clave] is None]
        if faltan_obligatorias:
            raise FinderError(
                "No pude identificar las columnas de nombre y teléfono en la planilla.")

        filas = []
        for cruda in crudas[1:]:
            nombre = _valor(cruda, indices["nombre"])
            if not nombre:
                continue
            fila = {"nombre": nombre,
                   "telefono": _formatear_telefono(_valor(cruda, indices["telefono"]))}
            for clave, _, obligatoria in _COLUMNAS:
                if obligatoria:
                    continue
                valor = _valor(cruda, indices[clave])
                if valor:
                    fila[clave] = valor
            filas.append(fila)

    cache.set(CACHE_FILAS, filas, settings.AZERTA_FINDER_CACHE_TTL)
    return filas


def _tokens_busqueda(texto):
    limpio = "".join(c if c.isalnum() else " " for c in normalizar(texto or ""))
    return [t for t in limpio.split() if len(t) >= 3]


def _texto_de_fila(fila):
    """Nombre + cargo + organizacion: lo que se busca, no el mail ni el
    telefono (buscar "56999813647" no es el caso de uso)."""
    return " ".join(t for t in (fila.get("nombre"), fila.get("cargo"),
                                fila.get("organizacion")) if t)


def buscar(consulta):
    """Contactos que calzan con `consulta`: nombre, cargo U organizacion -no
    hace falta el nombre completo, ni que la respuesta sea una sola persona.

    Se queda con las filas que comparten la MAYOR cantidad de palabras de la
    consulta (no cualquiera que comparta una sola): "gerente general de
    Viña Santa Rita" no debe traer a cualquier otro gerente general, y
    "contactos de Amchan" trae a TODOS los de Amchan -a todos les toca el
    mismo puntaje ("amchan"), asi que quedan todos, no solo el primero.
    Lista vacia si ninguna fila comparte ni una palabra.
    """
    tokens = _tokens_busqueda(consulta)
    if not tokens:
        return []

    puntuadas = []
    for fila in cargar():
        tokens_fila = set(_tokens_busqueda(_texto_de_fila(fila)))
        puntaje = sum(1 for t in tokens if t in tokens_fila)
        if puntaje:
            puntuadas.append((puntaje, fila))

    if not puntuadas:
        return []
    mejor = max(puntaje for puntaje, _ in puntuadas)
    return [fila for puntaje, fila in puntuadas if puntaje == mejor]
